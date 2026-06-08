# -*- coding: utf-8 -*-
import csv
import json
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

sys.path.append(str(Path(__file__).resolve().parents[1]))
from funcoes_auxiliares.padronizacao_csv import salvar_csv_padronizado

arquivos_entrada = [
    {
        'rotulo': 'TIPO 1 A 3',
        'arquivo': Path('data_exec_indiv/avaliacoes/11_2_base_tipo_1_a_3.csv'),
        'aba': 'TIPO_1_A_3',
    },
    {
        'rotulo': 'TIPO 4 A 7',
        'arquivo': Path('data_exec_indiv/avaliacoes/11_2_base_tipo_4_a_7.csv'),
        'aba': 'TIPO_4_A_7',
    },
    {
        'rotulo': 'TIPO 8 OU MAIS',
        'arquivo': Path('data_exec_indiv/avaliacoes/11_2_base_tipo_8_ou_mais.csv'),
        'aba': 'TIPO_8_OU_MAIS',
    },
]

arquivo_saida_excel = Path('data_exec_indiv/avaliacoes/12_resultado_final.xlsx')

pasta_resumo = Path('saida_resumo_avaliacoes') / 'exec_12_gerar_excel_final'
arquivo_resumo_json = pasta_resumo / 'exec_12_gerar_excel_final_resumo.json'
arquivo_resumo_txt = pasta_resumo / 'exec_12_gerar_excel_final_resumo.txt'
arquivo_resumo_csv = pasta_resumo / 'exec_12_gerar_excel_final_resumo.csv'

limite_linhas_excel = 1_048_576
limite_colunas_excel = 16_384

colunas_numericas = {
    'NOTA1',
    'NOTA2',
    'NOTA3',
    'NOTA4',
    'NOTA5',
    'NOTA GERAL',
    'TIPO',
    'META',
    'RESULTADO DA UNIDADE',
    'Quantidade notas validas',
    'MES',
    'DIA',
    'ANO',
}

# Identificadores e datas ficam como texto para o Excel nao remover zeros,
# nao converter codigos tipo "09E9" em notacao cientifica e nao reinterpretar datas.
colunas_texto_forcado = {
    'CDATENDIMENTO',
    'NUM_BENEFICIARIO',
    'CDEMPRESA',
    'CDUSUARIO',
    'TELEFONE',
    'DATA_ATENDIMENTO',
    'DT_RESPOSTA',
    'EMAIL',
    'NOME',
}


def valor_vazio(valor):
    return valor is None or str(valor).strip() == ''


def converter_numero(valor):
    if valor_vazio(valor):
        return None

    texto = str(valor).strip().replace(',', '.')

    try:
        numero = Decimal(texto)
    except InvalidOperation:
        return str(valor)

    if numero == numero.to_integral_value():
        return int(numero)

    return float(numero)


def criar_celula_texto(ws, valor):
    celula = WriteOnlyCell(ws, value='' if valor is None else str(valor))
    celula.number_format = '@'
    return celula


def criar_celula(ws, coluna, valor):
    if coluna in colunas_texto_forcado:
        return criar_celula_texto(ws, valor)

    if coluna in colunas_numericas:
        return converter_numero(valor)

    if valor_vazio(valor):
        return None

    return str(valor)


def validar_arquivos_obrigatorios():
    return [item['arquivo'] for item in arquivos_entrada if not item['arquivo'].exists()]


def contar_linhas_colunas_csv(caminho):
    with open(caminho, 'r', encoding='utf-8-sig', newline='') as arquivo:
        leitor = csv.reader(arquivo)
        cabecalho = next(leitor)
        total_linhas = sum(1 for _ in leitor)

    return total_linhas, len(cabecalho), cabecalho


def validar_limites_excel():
    problemas = []

    for item in arquivos_entrada:
        total_linhas, total_colunas, _ = contar_linhas_colunas_csv(item['arquivo'])
        linhas_com_cabecalho = total_linhas + 1

        if linhas_com_cabecalho > limite_linhas_excel:
            problemas.append(
                f"{item['arquivo']} possui {linhas_com_cabecalho} linhas com cabecalho "
                f"e excede o limite do Excel ({limite_linhas_excel})."
            )

        if total_colunas > limite_colunas_excel:
            problemas.append(
                f"{item['arquivo']} possui {total_colunas} colunas "
                f"e excede o limite do Excel ({limite_colunas_excel})."
            )

    return problemas


def escrever_aba_csv(workbook, item):
    caminho = item['arquivo']
    ws = workbook.create_sheet(title=item['aba'])

    with open(caminho, 'r', encoding='utf-8-sig', newline='') as arquivo:
        leitor = csv.reader(arquivo)
        cabecalho = next(leitor)
        ws.append(cabecalho)

        total_linhas = 0
        for linha in leitor:
            ws.append([
                criar_celula(ws, coluna, valor)
                for coluna, valor in zip(cabecalho, linha)
            ])
            total_linhas += 1

    return {
        'rotulo': item['rotulo'],
        'aba': item['aba'],
        'arquivo_csv': str(caminho),
        'linhas': total_linhas,
        'colunas': len(cabecalho),
    }


def escrever_aba_resumo(workbook, registros):
    ws = workbook.create_sheet(title='RESUMO', index=0)
    cabecalho = ['ROTULO', 'ABA', 'ARQUIVO_CSV', 'LINHAS', 'COLUNAS']
    ws.append(cabecalho)

    for registro in registros:
        ws.append([
            registro['rotulo'],
            registro['aba'],
            registro['arquivo_csv'],
            registro['linhas'],
            registro['colunas'],
        ])


def gerar_excel():
    workbook = Workbook(write_only=True)
    registros = []

    for item in arquivos_entrada:
        print(f"Gravando aba {item['aba']}: {item['arquivo']}")
        registros.append(escrever_aba_csv(workbook, item))

    escrever_aba_resumo(workbook, registros)
    arquivo_saida_excel.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(arquivo_saida_excel)
    return registros


def salvar_resumos(registros):
    pasta_resumo.mkdir(parents=True, exist_ok=True)

    resumo = {
        'execucao': 'exec_12_gerar_excel_final',
        'arquivo_saida_excel': str(arquivo_saida_excel),
        'abas': registros,
        'observacao': (
            'Colunas identificadoras e datas foram gravadas como texto para '
            'evitar perda de zeros a esquerda e conversao automatica do Excel.'
        ),
        'colunas_texto_forcado': sorted(colunas_texto_forcado),
        'colunas_numericas': sorted(colunas_numericas),
    }

    with open(arquivo_resumo_json, 'w', encoding='utf-8') as arquivo:
        json.dump(resumo, arquivo, ensure_ascii=False, indent=4)

    linhas_txt = [
        'RESUMO DA EXECUCAO 12 - GERAR EXCEL FINAL',
        '',
        f"Arquivo Excel: {resumo['arquivo_saida_excel']}",
        '',
        'ABAS GERADAS:',
    ]

    for registro in registros:
        linhas_txt.append(
            f"- {registro['aba']}: {registro['linhas']} linha(s), "
            f"{registro['colunas']} coluna(s)"
        )

    linhas_txt.extend([
        '',
        'Colunas identificadoras e datas gravadas como texto.',
    ])

    with open(arquivo_resumo_txt, 'w', encoding='utf-8') as arquivo:
        arquivo.write('\n'.join(linhas_txt))

    salvar_csv_padronizado(pd.DataFrame(registros), arquivo_resumo_csv)


def executar():
    print('Iniciando execucao 12 - gerar Excel final...')

    arquivos_faltando = validar_arquivos_obrigatorios()
    if arquivos_faltando:
        print('ERRO - arquivos obrigatorios nao encontrados:')
        for arquivo in arquivos_faltando:
            print(f'- {arquivo}')
        return 1

    problemas_limite = validar_limites_excel()
    if problemas_limite:
        print('ERRO - arquivos excedem limite do Excel:')
        for problema in problemas_limite:
            print(f'- {problema}')
        return 1

    registros = gerar_excel()
    salvar_resumos(registros)

    print(f'Excel final gerado: {arquivo_saida_excel}')
    print('Execucao 12 finalizada.')
    return 0


if __name__ == '__main__':
    sys.exit(executar())
